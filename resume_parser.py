"""Extract text from PDF, DOCX and TXT resumes. Works fully in memory:
nothing is written to disk, so no resume is stored."""
import io
from pathlib import Path

from pypdf import PdfReader
from docx import Document

MAX_SIZE_MB = 5
ALLOWED = {".pdf", ".docx", ".txt"}


class ResumeParseError(Exception):
    pass


def validate(filename: str, data: bytes) -> None:
    ext = Path(filename).suffix.lower()
    if ext not in ALLOWED:
        raise ResumeParseError(f"Unsupported file type '{ext}'. Use PDF, DOCX or TXT.")
    if len(data) > MAX_SIZE_MB * 1024 * 1024:
        raise ResumeParseError(f"File is larger than {MAX_SIZE_MB} MB.")
    if len(data) == 0:
        raise ResumeParseError("File is empty.")


def _pdf(data: bytes) -> str:
    try:
        reader = PdfReader(io.BytesIO(data))
        if reader.is_encrypted:
            raise ResumeParseError("PDF is password protected.")
        return "\n".join((page.extract_text() or "") for page in reader.pages)
    except ResumeParseError:
        raise
    except Exception as exc:
        raise ResumeParseError(f"Could not read PDF: {exc}") from exc


def _docx(data: bytes) -> str:
    try:
        doc = Document(io.BytesIO(data))
    except Exception as exc:
        raise ResumeParseError(f"Could not read DOCX: {exc}") from exc
    parts = [p.text for p in doc.paragraphs]
    for table in doc.tables:  # many resume templates put skills in tables
        for row in table.rows:
            for cell in row.cells:
                parts.append(cell.text)
    return "\n".join(parts)


def extract_text(filename: str, data: bytes) -> str:
    validate(filename, data)
    ext = Path(filename).suffix.lower()
    if ext == ".pdf":
        text = _pdf(data)
    elif ext == ".docx":
        text = _docx(data)
    else:
        text = data.decode("utf-8", errors="ignore")
    if len(text.strip()) < 30:
        raise ResumeParseError(
            "Almost no text was found. The resume may be a scanned image (OCR is not included)."
        )
    return text
